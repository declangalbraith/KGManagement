import logging
import zipfile
from hashlib import md5
from io import BytesIO
from datetime import datetime
from time import sleep

from django.utils.encoding import force_str
from django.utils.translation import gettext_lazy as _
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from django.core.files.base import ContentFile

from application.celery import app
from dvadmin.system.models import DownloadCenter, Users
from dvadmin.utils.import_export import build_import_request, execute_import_by_view
from dvadmin.utils.validator import CustomValidationError

logger = logging.getLogger(__name__)


def _import_failure_user_message(exc: BaseException) -> str:
    if isinstance(exc, CustomValidationError):
        return force_str(
            _(
                "The import could not be completed. Some data in the file is invalid. "
                "Please check your entries against the template, then try again."
            )
        )
    if isinstance(exc, (zipfile.BadZipFile, OSError)):
        return force_str(
            _(
                "The import could not be completed. The file could not be read. "
                "Please confirm the file is a valid Excel file and upload it again."
            )
        )
    return force_str(
        _(
            "The import could not be completed. Please verify the file format and your data, then try again. "
            "If the issue persists, contact the administrator."
        )
    )


def is_number(num):
    try:
        float(num)
        return True
    except ValueError:
        pass

    try:
        import unicodedata
        unicodedata.numeric(num)
        return True
    except (TypeError, ValueError):
        pass
    return False

def get_string_len(string):
    """
    获取字符串最大长度
    :param string:
    :return:
    """
    length = 4
    if string is None:
        return length
    if is_number(string):
        return length
    for char in string:
        length += 2.1 if ord(char) > 256 else 1
    return round(length, 1) if length <= 50 else 50

@app.task
def async_export_data(data: list, filename: str, dcid: int, export_field_label: dict):
    instance = DownloadCenter.objects.get(pk=dcid)
    instance.task_status = 1
    instance.save()
    sleep(2)
    try:
        wb = Workbook()
        ws = wb.active
        header_data = ["序号", *export_field_label.values()]
        hidden_header = ["#", *export_field_label.keys()]
        df_len_max = [get_string_len(ele) for ele in header_data]
        row = get_column_letter(len(export_field_label) + 1)
        column = 1
        ws.append(header_data)
        for index, results in enumerate(data):
            results_list = []
            for h_index, h_item in enumerate(hidden_header):
                for key, val in results.items():
                    if key == h_item:
                        if val is None or val == "":
                            results_list.append("")
                        elif isinstance(val, datetime):
                            val = val.strftime("%Y-%m-%d %H:%M:%S")
                            results_list.append(val)
                        else:
                            results_list.append(val)
                        # 计算最大列宽度
                        result_column_width = get_string_len(val)
                        if h_index != 0 and result_column_width > df_len_max[h_index]:
                            df_len_max[h_index] = result_column_width
            ws.append([index + 1, *results_list])
            column += 1
        # 　更新列宽
        for index, width in enumerate(df_len_max):
            ws.column_dimensions[get_column_letter(index + 1)].width = width
        tab = Table(displayName="Table", ref=f"A1:{row}{column}")  # 名称管理器
        style = TableStyleInfo(
            name="TableStyleLight11",
            showFirstColumn=True,
            showLastColumn=True,
            showRowStripes=True,
            showColumnStripes=True,
        )
        tab.tableStyleInfo = style
        ws.add_table(tab)
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        s = md5()
        while True:
            chunk = stream.read(1024)
            if not chunk:
                break
            s.update(chunk)
        stream.seek(0)
        instance.md5sum = s.hexdigest()
        instance.file_name = filename
        instance.url.save(filename, ContentFile(stream.read()))
        instance.task_status = 2
    except Exception as e:
        instance.task_status = 3
        instance.description = str(e)[:250]
    instance.save()


def send_import_message(user_id, title, content):
    user = Users.objects.filter(pk=user_id).first()
    if user is None:
        return
    # Delay the import to avoid a startup-time circular import:
    # tasks -> websocketConfig -> message_center -> viewset -> import_export_mixin -> tasks
    from application.websocketConfig import create_message_push

    create_message_push(
        title=title,
        content=content,
        target_user=[user_id],
        message={
            "sender": "system",
            "contentType": "SYSTEM",
            "content": content,
        },
        request=build_import_request(user=user, path="/api/system/message_center/"),
    )


def try_send_import_message(user_id, title, content):
    try:
        send_import_message(user_id=user_id, title=title, content=content)
    except Exception:
        logger.exception("Failed to send import notification")


@app.task
def async_import_data(viewset_path: str, user_id: int, file_url: str, request_path: str, import_title: str,
                      periodic_task_name: str = None):
    try:
        row_count = execute_import_by_view(
            viewset_path=viewset_path,
            user_id=user_id,
            file_url=file_url,
            request_path=request_path,
        )
        try_send_import_message(
            user_id=user_id,
            title=import_title,
            content=_("Import completed successfully. %(count)s rows were processed.") % {"count": row_count},
        )
        return {"status": "success", "row_count": row_count, "periodic_task_name": periodic_task_name}
    except Exception as exc:
        try_send_import_message(
            user_id=user_id,
            title=import_title,
            content=_import_failure_user_message(exc),
        )
        raise
