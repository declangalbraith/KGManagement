import csv
import io
from dataclasses import dataclass

from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

REQUIRED_COLUMNS = ("Number", "State", "Type Designation", "Description1 (EN)")
BOM_LEVEL_COLUMN = "BOM Level"
ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".xls"}


@dataclass
class BomRootFields:
    number: str
    state: str
    type_designation: str
    description_en: str


def _normalize_header(name: str) -> str:
    if name is None:
        return ""
    return str(name).strip().lstrip("\ufeff")


def _validate_columns(fieldnames: list[str]) -> None:
    normalized = {_normalize_header(f) for f in fieldnames if f}
    missing = [col for col in REQUIRED_COLUMNS if col not in normalized]
    if missing:
        raise ValidationError(
            {"file": f"BOM 文件缺少必要列: {', '.join(missing)}"}
        )


def _row_dict(raw_row: dict) -> dict[str, str]:
    result: dict[str, str] = {}
    for key, value in raw_row.items():
        norm_key = _normalize_header(key)
        if not norm_key:
            continue
        if value is None:
            result[norm_key] = ""
        else:
            result[norm_key] = str(value).strip()
    return result


def _extract_root_from_rows(rows: list[dict]) -> BomRootFields:
    if not rows:
        raise ValidationError({"file": "BOM 文件没有数据行"})

    root_row = None
    for row in rows:
        data = _row_dict(row)
        level = data.get(BOM_LEVEL_COLUMN, "")
        if level in ("0", "0.0"):
            root_row = data
            break

    if root_row is None:
        root_row = _row_dict(rows[0])

    number = root_row.get("Number", "").strip()
    if not number:
        raise ValidationError({"file": "BOM 根节点 Number 为空，无法识别文档"})

    return BomRootFields(
        number=number,
        state=root_row.get("State", ""),
        type_designation=root_row.get("Type Designation", ""),
        description_en=root_row.get("Description1 (EN)", ""),
    )


def parse_csv_bytes(content: bytes) -> BomRootFields:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValidationError({"file": "BOM CSV 文件缺少表头"})
    _validate_columns(list(reader.fieldnames))
    rows = list(reader)
    return _extract_root_from_rows(rows)


def parse_xlsx_bytes(content: bytes) -> BomRootFields:
    wb = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            raise ValidationError({"file": "BOM Excel 文件缺少表头"})
        fieldnames = [_normalize_header(str(h) if h is not None else "") for h in header]
        _validate_columns(fieldnames)
        dict_rows = []
        for row in rows_iter:
            if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
                continue
            item = {}
            for idx, col in enumerate(fieldnames):
                if not col:
                    continue
                value = row[idx] if idx < len(row) else ""
                item[col] = "" if value is None else value
            dict_rows.append(item)
        return _extract_root_from_rows(dict_rows)
    finally:
        wb.close()


def parse_bom_file(filename: str, content: bytes) -> BomRootFields:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return parse_csv_bytes(content)
    if lower.endswith(".xlsx"):
        return parse_xlsx_bytes(content)
    if lower.endswith(".xls"):
        raise ValidationError({"file": "暂不支持 .xls 格式，请使用 .xlsx 或 .csv"})
    raise ValidationError({"file": "仅支持 Excel（.xlsx）或 CSV（.csv）文件"})


def build_minio_object_key(number: str, filename: str) -> str:
    safe_number = number.strip().replace("/", "_").replace("\\", "_")
    safe_name = filename.strip().replace("/", "_").replace("\\", "_")
    return f"docManage/BOM-Management/{safe_number}/{safe_name}"
