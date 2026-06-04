"""Extract general document body for in-browser preview (PDF stream / Word HTML)."""

from __future__ import annotations

import html
import io

PREVIEWABLE_PDF = frozenset({"pdf"})
PREVIEWABLE_WORD = frozenset({"docx"})


def normalize_ext(file_ext: str) -> str:
    return (file_ext or "").lower().lstrip(".")


def docx_bytes_to_html(content: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(f"<p>{html.escape(text)}</p>")
    for table in doc.tables:
        rows_html: list[str] = []
        for row in table.rows:
            cells = "".join(
                f"<td>{html.escape(cell.text.strip())}</td>" for cell in row.cells
            )
            rows_html.append(f"<tr>{cells}</tr>")
        if rows_html:
            parts.append(f"<table>{''.join(rows_html)}</table>")
    if not parts:
        raise ValueError("docx 文件未提取到可预览的正文内容")
    return "\n".join(parts)
