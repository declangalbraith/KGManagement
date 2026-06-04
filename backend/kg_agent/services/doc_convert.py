"""Convert legacy Word .doc to .docx for 8D pipeline (facade accepts docx/pdf only)."""

from __future__ import annotations

import logging
import os
import shutil
import subprocess
import sys
import tempfile
from typing import List, Optional

logger = logging.getLogger(__name__)

_SOFFICE_CANDIDATES = (
    "soffice",
    "libreoffice",
    "loffice",
)
_WINDOWS_SOFFICE_PATHS = (
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
)
_TEXT_EXTRACTORS = (
    (["antiword", "-m", "UTF-8"], "antiword"),
    (["catdoc", "-w"], "catdoc"),
)


def _resolve_soffice() -> Optional[str]:
    for name in _SOFFICE_CANDIDATES:
        found = shutil.which(name)
        if found:
            return found
    if sys.platform == "win32":
        for path in _WINDOWS_SOFFICE_PATHS:
            if os.path.isfile(path):
                return path
    return None


def _convert_with_soffice(src_path: str, dst_path: str) -> bool:
    soffice = _resolve_soffice()
    if not soffice:
        return False
    out_dir = os.path.dirname(dst_path) or "."
    os.makedirs(out_dir, exist_ok=True)
    cmd = [
        soffice,
        "--headless",
        "--norestore",
        "--convert-to",
        "docx",
        "--outdir",
        out_dir,
        src_path,
    ]
    try:
        subprocess.run(
            cmd,
            check=True,
            capture_output=True,
            timeout=180,
            text=True,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as exc:
        logger.warning("LibreOffice .doc→.docx failed: %s", exc)
        return False

    produced = os.path.join(
        out_dir, f"{os.path.splitext(os.path.basename(src_path))[0]}.docx"
    )
    if not os.path.isfile(produced):
        return False
    if os.path.abspath(produced) != os.path.abspath(dst_path):
        if os.path.isfile(dst_path):
            os.unlink(dst_path)
        shutil.move(produced, dst_path)
    return True


def _convert_with_word_com(src_path: str, dst_path: str) -> bool:
    if sys.platform != "win32":
        return False
    try:
        import win32com.client  # type: ignore[import-untyped]
    except ImportError:
        return False

    word = None
    doc = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(os.path.abspath(src_path), ReadOnly=True)
        doc.SaveAs2(os.path.abspath(dst_path), FileFormat=16)
        doc.Close(False)
        return os.path.isfile(dst_path) and os.path.getsize(dst_path) > 0
    except Exception as exc:
        logger.warning("MS Word COM .doc→.docx failed: %s", exc)
        return False
    finally:
        if doc is not None:
            try:
                doc.Close(False)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass


def _extract_doc_plain_text(src_path: str) -> Optional[str]:
    for base_args, _label in _TEXT_EXTRACTORS:
        exe = base_args[0]
        if not shutil.which(exe):
            continue
        cmd: List[str] = [*base_args, src_path]
        try:
            proc = subprocess.run(
                cmd,
                check=True,
                capture_output=True,
                timeout=120,
                text=True,
                encoding="utf-8",
                errors="replace",
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError):
            continue
        text = (proc.stdout or "").strip()
        if text:
            return text
    return None


def _convert_via_text_docx(src_path: str, dst_path: str) -> bool:
    text = _extract_doc_plain_text(src_path)
    if not text:
        return False
    try:
        from docx import Document
    except ImportError:
        return False

    doc = Document()
    for block in text.splitlines():
        line = block.strip()
        if line:
            doc.add_paragraph(line)
        else:
            doc.add_paragraph("")
    if not doc.paragraphs:
        return False
    doc.save(dst_path)
    return os.path.isfile(dst_path) and os.path.getsize(dst_path) > 0


def convert_doc_to_docx(src_path: str) -> str:
    """
    Convert .doc to a new .docx file in a temp directory.
    Returns absolute path to the .docx file.
    """
    if not src_path.lower().endswith(".doc"):
        raise ValueError("convert_doc_to_docx 仅用于 .doc 文件")
    if not os.path.isfile(src_path):
        raise FileNotFoundError(f"文件不存在: {src_path}")

    out_dir = tempfile.mkdtemp(prefix="kg_doc_convert_")
    base = os.path.splitext(os.path.basename(src_path))[0]
    dst_path = os.path.join(out_dir, f"{base}.docx")

    for converter in (
        _convert_with_soffice,
        _convert_with_word_com,
        _convert_via_text_docx,
    ):
        if converter(src_path, dst_path):
            return dst_path

    raise ValueError(
        "无法将 .doc 转为 .docx。请任选其一："
        "① 在 Word 中另存为 .docx 后上传；"
        "② 在服务器安装 LibreOffice（soffice）；"
        "③ Windows 安装 Microsoft Word 并执行 pip install pywin32；"
        "④ Linux 安装 antiword/catdoc 以启用纯文本回退转换。"
    )
